"""Phase 11: uploaded documents -> searchable knowledge chunks.

    upload (pdf / docx / txt) -> check type, size and content -> extract text page by page
        -> cut into passages -> TF-IDF keywords per passage -> KnowledgeChunk rows (searched by Phase 7)

Only ACTIVE documents are searched (Phase 7 already filters on Document.status). A replaced or retired document
is archived, never deleted, so old versions can be restored.

pypdf and python-docx are imported only when a file of that type is processed, so the rest of the app still
starts if one of them is not installed.
"""
import io
import logging
import re
import zipfile
from dataclasses import dataclass
from pathlib import Path

from sqlalchemy import delete, select
from sqlalchemy.orm import Session

from app.core.config import ROOT, get_settings
from app.models import Document, KnowledgeChunk
from app.services.retrieval import split_passages
from app.services.text_mining import get_text_miner

logger = logging.getLogger("pgpc.documents")

STORAGE_DIR = ROOT / "data" / "documents"
ALLOWED_TYPES = {"pdf", "docx", "txt"}
MAX_CHUNKS = 3000
MAX_UNZIPPED_BYTES = 80 * 1024 * 1024      # a .docx is a zip; refuse "zip bombs"


class DocumentError(Exception):
    """A problem the uploader can fix. `status` is the HTTP status to answer with."""

    def __init__(self, message: str, status: int = 422):
        super().__init__(message)
        self.message, self.status = message, status


@dataclass
class Page:
    label: str | None      # "3" for PDF page 3, None when the format has no pages
    text: str


# --------------------------------------------------------------------------- file checks
def safe_name(name: str | None) -> str:
    """Keeps only the file's own name: no folders, no odd characters."""
    base = re.split(r"[\\/]", name or "")[-1]
    base = re.sub(r"[^A-Za-z0-9._ ()-]", "_", base).strip(" .")
    return (base or "document")[:200]


def file_type_of(name: str) -> str:
    ext = Path(name).suffix.lower().lstrip(".")
    if ext not in ALLOWED_TYPES:
        raise DocumentError("Only PDF, DOCX and TXT files can be uploaded.", 415)
    return ext


def check_content(data: bytes, ftype: str) -> None:
    """The file extension is only a claim, so the content must really look like that type."""
    if not data:
        raise DocumentError("The file is empty.")
    if ftype == "pdf" and not data.lstrip()[:5] == b"%PDF-":
        raise DocumentError("This file is not a real PDF.")
    if ftype == "docx":
        try:
            with zipfile.ZipFile(io.BytesIO(data)) as z:
                if "word/document.xml" not in z.namelist():
                    raise DocumentError("This file is not a real Word (.docx) document.")
                if sum(i.file_size for i in z.infolist()) > MAX_UNZIPPED_BYTES:
                    raise DocumentError("This Word document is too large once unpacked.")
        except zipfile.BadZipFile:
            raise DocumentError("This file is not a real Word (.docx) document.")
    if ftype == "txt" and b"\x00" in data[:4096]:
        raise DocumentError("This file is not a plain text file.")


# --------------------------------------------------------------------------- text extraction
def _decode(data: bytes) -> str:
    for enc in ("utf-8-sig", "utf-16", "cp1252"):
        try:
            return data.decode(enc)
        except (UnicodeDecodeError, UnicodeError):
            continue
    return data.decode("latin-1", errors="replace")


def _clean_block(text: str) -> str:
    text = text.replace("\x00", "").replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"[ \t\f\v]+", " ", text)
    return re.sub(r"\n{3,}", "\n\n", text).strip()


def extract_pages(data: bytes, ftype: str) -> list[Page]:
    if ftype == "txt":
        pages = [Page(None, _clean_block(_decode(data)))]
    elif ftype == "pdf":
        try:
            from pypdf import PdfReader
        except ImportError:
            raise DocumentError("PDF support is not installed on the server (pip install pypdf).", 501)
        try:
            reader = PdfReader(io.BytesIO(data))
            if reader.is_encrypted:
                raise DocumentError("This PDF is password protected. Remove the password and upload it again.")
            pages = [Page(str(i), _clean_block(p.extract_text() or "")) for i, p in enumerate(reader.pages, 1)]
        except DocumentError:
            raise
        except Exception:
            raise DocumentError("This PDF could not be read. It may be damaged.")
    else:  # docx
        try:
            from docx import Document as WordDocument
        except ImportError:
            raise DocumentError("Word support is not installed on the server (pip install python-docx).", 501)
        try:
            doc = WordDocument(io.BytesIO(data))
            blocks = [p.text for p in doc.paragraphs if p.text.strip()]
            for table in doc.tables:
                for row in table.rows:
                    cells = [c.text.strip() for c in row.cells if c.text.strip()]
                    if cells:
                        blocks.append(" | ".join(cells))
            pages = [Page(None, _clean_block("\n\n".join(blocks)))]
        except Exception:
            raise DocumentError("This Word document could not be read. It may be damaged.")
    pages = [p for p in pages if p.text]
    if not pages:
        raise DocumentError("No readable text was found in this file. A scanned PDF (pictures of pages) "
                            "needs OCR before it can be used.")
    return pages


# --------------------------------------------------------------------------- chunking
@dataclass
class ChunkData:
    content: str
    page: str | None
    keywords: str


def make_chunks(pages: list[Page]) -> list[ChunkData]:
    pieces: list[tuple[str, str | None]] = []
    for page in pages:
        for passage in split_passages(page.text):
            if len(passage) >= 20:               # skip page numbers and stray headers
                pieces.append((passage, page.label))
    if not pieces:
        raise DocumentError("No readable text was found in this file.")
    if len(pieces) > MAX_CHUNKS:
        raise DocumentError(f"This document is too long (over {MAX_CHUNKS} passages). Split it into parts.")
    miner = get_text_miner()
    keywords = miner.tfidf_keywords([p for p, _ in pieces], top_n=8) if len(pieces) > 1 else \
        [miner.document_keywords(pieces[0][0], top_n=8)]
    return [ChunkData(text, page, ", ".join(kw)) for (text, page), kw in zip(pieces, keywords)]


# --------------------------------------------------------------------------- storage
def stored_path(document_id: int, ftype: str) -> Path:
    return STORAGE_DIR / f"{document_id}.{ftype}"


def write_file(document_id: int, ftype: str, data: bytes) -> Path:
    STORAGE_DIR.mkdir(parents=True, exist_ok=True)
    path = stored_path(document_id, ftype)
    path.write_bytes(data)
    return path


def remove_file(document_id: int, ftype: str) -> None:
    try:
        stored_path(document_id, ftype).unlink(missing_ok=True)
    except OSError:
        logger.warning("could not remove stored file for document id=%s", document_id)


# --------------------------------------------------------------------------- database work
def chunk_source(doc: Document) -> str:
    return (doc.source or doc.title)[:200]


def replace_chunks(db: Session, doc: Document, chunks: list[ChunkData], category: str | None) -> None:
    db.execute(delete(KnowledgeChunk).where(KnowledgeChunk.document_id == doc.id))
    for i, c in enumerate(chunks):
        db.add(KnowledgeChunk(document_id=doc.id, category=category, title=doc.title[:300], content=c.content,
                              keywords=c.keywords, chunk_index=i, source=chunk_source(doc), source_page=c.page))


def title_in_use(db: Session, title: str, exclude_id: int | None = None) -> bool:
    q = select(Document.id).where(Document.status == "active", Document.title.ilike(title))
    if exclude_id is not None:
        q = q.where(Document.id != exclude_id)
    return db.scalar(q) is not None


def create_document(db: Session, *, data: bytes, original_name: str, title: str | None, description: str | None,
                    source: str | None, category: str | None, user_id: int | None,
                    replaces_id: int | None = None) -> Document:
    """Validates, extracts, stores. Nothing is saved unless every step works."""
    name = safe_name(original_name)
    ftype = file_type_of(name)
    limit = get_settings().max_upload_size_mb * 1024 * 1024
    if len(data) > limit:
        raise DocumentError(f"The file is larger than {get_settings().max_upload_size_mb} MB.", 413)
    check_content(data, ftype)
    chunks = make_chunks(extract_pages(data, ftype))

    title = (title or Path(name).stem).strip()[:300]
    old = None
    if replaces_id is not None:
        old = db.get(Document, replaces_id)
        if old is None:
            raise DocumentError("The document to replace was not found.", 404)
    if title_in_use(db, title, exclude_id=old.id if old else None):
        raise DocumentError("An active document with that title already exists. Use a different title, or "
                            "replace the existing one.", 409)

    doc = Document(title=title, file_name=name, file_type=ftype, description=description, source=source,
                   uploaded_by=user_id, status="active")
    db.add(doc)
    db.flush()
    path = None
    try:
        path = write_file(doc.id, ftype, data)
        replace_chunks(db, doc, chunks, category)
        if old is not None:
            old.status = "archived"          # the previous version stays on file, just no longer searched
        db.commit()
    except Exception:
        db.rollback()
        if path is not None:
            remove_file(doc.id, ftype)
        raise
    db.refresh(doc)
    logger.info("document uploaded id=%s type=%s chunks=%d replaces=%s", doc.id, ftype, len(chunks), replaces_id)
    return doc


def reprocess(db: Session, doc: Document) -> int:
    """Re-reads the stored file and rebuilds its chunks (for example after the text-mining rules improved)."""
    path = stored_path(doc.id, doc.file_type)
    if not path.exists():
        raise DocumentError("The stored file is missing, so it cannot be processed again.", 404)
    chunks = make_chunks(extract_pages(path.read_bytes(), doc.file_type))
    old = db.scalar(select(KnowledgeChunk.category).where(KnowledgeChunk.document_id == doc.id).limit(1))
    replace_chunks(db, doc, chunks, old)
    db.commit()
    return len(chunks)
